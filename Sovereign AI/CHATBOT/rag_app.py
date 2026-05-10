import streamlit as st
import requests
import time
import os
import io
from PyPDF2 import PdfReader
from docx import Document
import pandas as pd
from pptx import Presentation
import xml.etree.ElementTree as ET
from sentence_transformers import SentenceTransformer
import faiss
import numpy as np
from typing import List, Dict

st.set_page_config(
    page_title="RAG Llama 3.2 Chatbot",
    page_icon="",
    layout="wide"
)

st.title(" RAG Llama 3.2 Chatbot")

st.markdown("""
<style>
.user-message {
    background: linear-gradient(135deg, #0066ff 0%, #00ccff 100%);
    color: white;
    padding: 15px 20px;
    border-radius: 20px;
    margin: 10px 0;
    max-width: 80%;
    margin-left: auto;
    box-shadow: 0 4px 15px rgba(0, 102, 255, 0.6);
    animation: glow 2s ease-in-out infinite alternate;
}

.bot-message {
    background: linear-gradient(135deg, #11998e 0%, #38ef7d 100%);
    color: white;
    padding: 15px 20px;
    border-radius: 20px;
    margin: 10px 0;
    max-width: 80%;
    box-shadow: 0 4px 15px rgba(17, 153, 142, 0.4);
    animation: glow 2s ease-in-out infinite alternate;
}

@keyframes glow {
    from {
        box-shadow: 0 4px 15px rgba(0, 102, 255, 0.6);
    }
    to {
        box-shadow: 0 4px 25px rgba(0, 102, 255, 0.9);
    }
}

.loading-message {
    background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
    color: white;
    padding: 15px 20px;
    border-radius: 20px;
    margin: 10px 0;
    max-width: 80%;
    box-shadow: 0 4px 15px rgba(240, 147, 251, 0.4);
    animation: pulse 1.5s ease-in-out infinite;
}

@keyframes pulse {
    0% {
        opacity: 0.7;
        transform: scale(1);
    }
    50% {
        opacity: 1;
        transform: scale(1.02);
    }
    100% {
        opacity: 0.7;
        transform: scale(1);
    }
}

.chat-container {
    height: 70vh;
    overflow-y: auto;
    padding: 20px;
    background: #0e1117;
    border-radius: 15px;
    margin-bottom: 20px;
}

.stTextInput > div > div > input {
    background: white;
    border: 2px solid #667eea;
    border-radius: 25px;
    padding: 12px 20px;
    font-size: 16px;
}

.stTextInput > div > div > input:focus {
    border-color: #764ba2;
    box-shadow: 0 0 0 3px rgba(102, 126, 234, 0.2);
}

.sidebar {
    background: #1e1e2e;
    padding: 20px;
    border-radius: 15px;
    margin-bottom: 20px;
}

.document-item {
    background: #2a2a3e;
    padding: 10px;
    margin: 5px 0;
    border-radius: 8px;
    display: flex;
    justify-content: space-between;
    align-items: center;
}

.upload-area {
    border: 2px dashed #667eea;
    border-radius: 10px;
    padding: 20px;
    text-align: center;
    margin-bottom: 20px;
}
</style>
""", unsafe_allow_html=True)

# Initialize session state
if "messages" not in st.session_state:
    st.session_state.messages = []
if "documents" not in st.session_state:
    st.session_state.documents = {}
if "embeddings" not in st.session_state:
    st.session_state.embeddings = None
if "faiss_index" not in st.session_state:
    st.session_state.faiss_index = None
if "document_chunks" not in st.session_state:
    st.session_state.document_chunks = []

# Initialize embedding model
@st.cache_resource
def load_embedding_model():
    return SentenceTransformer('all-MiniLM-L6-v2', device='cpu')

embedding_model = load_embedding_model()

def parse_pdf(file) -> str:
    pdf_reader = PdfReader(io.BytesIO(file.read()))
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text() + "\n"
    return text

def parse_docx(file) -> str:
    doc = Document(io.BytesIO(file.read()))
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def parse_csv(file) -> str:
    df = pd.read_csv(io.BytesIO(file.read()))
    return df.to_string()

def parse_txt(file) -> str:
    return file.read().decode('utf-8')

def parse_xml(file) -> str:
    tree = ET.parse(io.BytesIO(file.read()))
    root = tree.getroot()
    text = ""
    for elem in root.iter():
        if elem.text:
            text += elem.text + " "
    return text

def parse_pptx(file) -> str:
    prs = Presentation(io.BytesIO(file.read()))
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def chunk_text(text: str, chunk_size: int = 500) -> List[str]:
    words = text.split()
    chunks = []
    for i in range(0, len(words), chunk_size):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
    return chunks

def create_embeddings(chunks: List[str]):
    embeddings = embedding_model.encode(chunks)
    dimension = embeddings.shape[1]
    index = faiss.IndexFlatL2(dimension)
    index.add(embeddings)
    return embeddings, index

def search_documents(query: str, k: int = 3) -> List[str]:
    if st.session_state.faiss_index is None:
        return []
    
    query_embedding = embedding_model.encode([query])
    distances, indices = st.session_state.faiss_index.search(query_embedding, k)
    
    relevant_chunks = []
    for idx in indices[0]:
        if idx < len(st.session_state.document_chunks):
            relevant_chunks.append(st.session_state.document_chunks[idx])
    
    return relevant_chunks

def get_rag_response(prompt: str) -> str:
    try:
        # Search for relevant documents
        relevant_docs = search_documents(prompt)
        
        # Create context from documents
        context = ""
        if relevant_docs:
            context = "\n\n".join(relevant_docs[:3])  # Use top 3 relevant chunks
            enhanced_prompt = f"""Based on the following context, please answer the user's question. If the context doesn't contain relevant information, say so and provide a general response.

Context:
{context}

User Question: {prompt}

Answer:"""
        else:
            enhanced_prompt = prompt
        
        response = requests.post(
            'http://localhost:11434/api/generate',
            json={
                'model': 'llama3.2:1b',
                'prompt': enhanced_prompt,
                'stream': False
            },
            timeout=30
        )
        
        if response.status_code == 200:
            return response.json()['response']
        else:
            return f"Error: Could not connect to Llama 3.2. Make sure Ollama is running with 'ollama serve'"
            
    except requests.exceptions.ConnectionError:
        return "Error: Cannot connect to Ollama. Please make sure Ollama is installed and running with 'ollama serve'. Install from https://ollama.ai/"
    except Exception as e:
        return f"Error: {str(e)}"

def get_llama_response(prompt: str) -> str:
    try:
        response = requests.post(
            'http://localhost:11434/api/generate',
            json={
                'model': 'llama3.2:1b',
                'prompt': prompt,
                'stream': False
            },
            timeout=30
        )
        
        if response.status_code == 200:
            return response.json()['response']
        else:
            return f"Error: Could not connect to Llama 3.2. Make sure Ollama is running with 'ollama serve'"
            
    except requests.exceptions.ConnectionError:
        return "Error: Cannot connect to Ollama. Please make sure Ollama is installed and running with 'ollama serve'. Install from https://ollama.ai/"
    except Exception as e:
        return f"Error: {str(e)}"

# Main layout
st.markdown("""
<style>
.main-container {
    display: flex;
    gap: 20px;
    height: 85vh;
}
.chat-section {
    flex: 3;
    display: flex;
    flex-direction: column;
}
.sidebar-section {
    flex: 1;
    background: #1e1e2e;
    padding: 20px;
    border-radius: 15px;
    overflow-y: auto;
}
.chat-messages {
    flex: 1;
    overflow-y: auto;
    padding: 20px;
    background: #0e1117;
    border-radius: 15px;
    margin-bottom: 15px;
}
.chat-input {
    background: #1e1e2e;
    padding: 15px;
    border-radius: 15px;
}
</style>
""", unsafe_allow_html=True)

# Main content area
with st.container():
    # Chat messages container
    with st.container():
        st.markdown('<div class="chat-messages">', unsafe_allow_html=True)
        
        # Display all messages
        for message in st.session_state.messages:
            if message["role"] == "user":
                st.markdown(f'<div class="user-message">👤 You: {message["content"]}</div>', unsafe_allow_html=True)
            else:
                st.markdown(f'<div class="bot-message">🤖 Bot: {message["content"]}</div>', unsafe_allow_html=True)
        
        st.markdown('</div>', unsafe_allow_html=True)
    
    # Chat input at bottom
    user_input = st.chat_input("Type your message here...")

    if user_input:
        # Add user message to session state
        st.session_state.messages.append({"role": "user", "content": user_input})
        
        # Show user message immediately
        st.markdown(f'<div class="user-message">👤 You: {user_input}</div>', unsafe_allow_html=True)
        
        # Show loading message
        loading_placeholder = st.empty()
        loading_placeholder.markdown('<div class="loading-message">🤖 ANSWERING TO THE QUERY...</div>', unsafe_allow_html=True)
        
        # Get response based on mode
        if st.session_state.get("use_rag", False) and st.session_state.faiss_index is not None:
            response = get_rag_response(user_input)
        else:
            response = get_llama_response(user_input)
        
        # Remove loading message and show response
        loading_placeholder.empty()
        
        # Add mode indicator to response
        mode_indicator = "📚 [RAG Mode] " if st.session_state.get("use_rag", False) and st.session_state.faiss_index is not None else "💬 [Chat Mode] "
        st.markdown(f'<div class="bot-message">🤖 {mode_indicator}Bot: {response}</div>', unsafe_allow_html=True)
        
        # Add bot response to session state
        st.session_state.messages.append({"role": "assistant", "content": response})
        
        # Rerun to update the interface
        st.rerun()

with st.sidebar:
    st.markdown('<div class="sidebar-section">', unsafe_allow_html=True)
    
    # Mode toggle
    st.subheader("🎯 Chat Mode")
    use_rag = st.toggle(
        "📚 Use RAG (Document-based)",
        value=st.session_state.get("use_rag", False),
        help="Toggle between RAG mode (uses documents) and Chat mode (normal responses)"
    )
    st.session_state["use_rag"] = use_rag
    
    if use_rag:
        st.info("📚 RAG Mode: Bot will use uploaded documents for context")
    else:
        st.info("💬 Chat Mode: Bot will provide general responses")
    
    st.divider()
    
    # Document upload
    st.subheader("📚 Document Upload")
    uploaded_files = st.file_uploader(
        "Upload Documents",
        type=['pdf', 'docx', 'csv', 'txt', 'xml', 'pptx'],
        accept_multiple_files=True,
        help="Upload PDF, Word, CSV, Text, XML, or PowerPoint files"
    )
    
    if uploaded_files:
        st.write(f"DEBUG: Found {len(uploaded_files)} files")
        documents_processed = False
        for file in uploaded_files:
            st.write(f"DEBUG: Processing file: {file.name}")
            file_extension = file.name.split('.')[-1].lower()
            
            try:
                if file_extension == 'pdf':
                    text = parse_pdf(file)
                elif file_extension == 'docx':
                    text = parse_docx(file)
                elif file_extension == 'csv':
                    text = parse_csv(file)
                elif file_extension == 'txt':
                    text = parse_txt(file)
                elif file_extension == 'xml':
                    text = parse_xml(file)
                elif file_extension == 'pptx':
                    text = parse_pptx(file)
                else:
                    st.error(f"Unsupported file type: {file_extension}")
                    continue
                
                # Store document
                st.session_state.documents[file.name] = text
                st.write(f"DEBUG: Stored document with {len(text)} characters")
                
                # Create chunks and embeddings
                chunks = chunk_text(text)
                st.session_state.document_chunks.extend(chunks)
                st.write(f"DEBUG: Created {len(chunks)} chunks")
                
                st.success(f" Uploaded: {file.name}")
                documents_processed = True
                
            except Exception as e:
                st.error(f"Error processing {file.name}: {str(e)}")
                st.write(f"DEBUG: Exception details: {e}")
        
        # Set flag for automatic processing
        if documents_processed:
            st.session_state.auto_process_pending = True
            st.write("DEBUG: Auto-process flag set")
    else:
        st.write("DEBUG: No files uploaded")
    
    # Automatic processing (triggered once)
    if st.session_state.get("auto_process_pending", False) and st.session_state.document_chunks:
        with st.spinner("🔄 Automatically creating embeddings..."):
            embeddings, index = create_embeddings(st.session_state.document_chunks)
            st.session_state.embeddings = embeddings
            st.session_state.faiss_index = index
        st.success("✅ Documents automatically processed and ready for RAG!")
        st.session_state.auto_process_pending = False  # Reset flag
    
    # Manual process button (backup option)
    if st.session_state.document_chunks and not st.session_state.get("auto_process_pending", False):
        if st.button("🔄 Process Documents"):
            with st.spinner("Creating embeddings..."):
                embeddings, index = create_embeddings(st.session_state.document_chunks)
                st.session_state.embeddings = embeddings
                st.session_state.faiss_index = index
            st.success("✅ Documents processed and ready for RAG!")
    
    st.divider()
    
    # Status
    st.subheader(" Status")
    
    st.write(f"DEBUG: Session state - Documents: {len(st.session_state.documents)}, Chunks: {len(st.session_state.document_chunks)}")
    
    if st.session_state.documents:
        st.success(f" {len(st.session_state.documents)} documents")
        st.success(f" {len(st.session_state.document_chunks)} chunks")
        if st.session_state.faiss_index is not None:
            st.success(" Embeddings ready")
            if use_rag:
                st.success(" RAG ACTIVE")
        else:
            st.warning(" Process documents")
        
        # Document list
        st.subheader(" Documents")
        for filename in list(st.session_state.documents.keys())[:3]:  # Show first 3
            with st.expander(f" {filename}"):
                st.text(st.session_state.documents[filename][:200] + "...")
        
        if st.button(" Clear All"):
            st.session_state.documents = {}
            st.session_state.embeddings = None
            st.session_state.faiss_index = None
            st.session_state.document_chunks = []
            st.rerun()
    else:
        st.info(" Upload documents")
    
    st.markdown('</div>', unsafe_allow_html=True)
